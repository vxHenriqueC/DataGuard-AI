# file_reader.py
import csv
import json
import os
import re
from io import StringIO, BytesIO
from typing import List, Tuple, Optional
import defusedxml.minidom as xml_minidom
from pypdf import PdfReader
from openpyxl import load_workbook
import docx

from settings import settings
from models import ExtractedUnit, FileExtractionResult

# Assinaturas de cabeçalho binário (Magic Bytes)
MAGIC_HEADERS = {
    "pdf": [b"%PDF-"],
    "xlsx": [b"PK\x03\x04"],
    "docx": [b"PK\x03\x04"],
    "xml": [b"<?xml", b"<"],
}

# Bloqueio estrito de extensões executáveis e perigosas
DANGEROUS_EXTENSIONS = {
    "exe", "bat", "cmd", "ps1", "sh", "vbs", "jar", "msi", "com", "scr", "dll", "so"
}


def sanitize_filename(filename: str) -> str:
    """Higieniza o nome do arquivo para prevenir Path Traversal."""
    clean_name = os.path.basename(filename)
    clean_name = re.sub(r'[^a-zA-Z0-9_.-]', '_', clean_name)
    return clean_name or "uploaded_file"


def validate_file_content(file_name: str, content: bytes) -> Tuple[bool, Optional[str]]:
    """Validação de segurança: tamanho, extensão e assinatura mágica."""
    file_size = len(content)

    if file_size == 0:
        return False, "Arquivo vazio."

    if file_size > settings.MAX_FILE_SIZE_MB * 1024 * 1024:
        return False, f"Arquivo excede o limite máximo permitido de {settings.MAX_FILE_SIZE_MB}MB."

    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""

    if ext in DANGEROUS_EXTENSIONS:
        return False, f"Arquivos com extensão .{ext} são bloqueados por segurança."

    if ext not in settings.ALLOWED_EXTENSIONS:
        return False, f"Formato .{ext} não suportado. Formatos aceitos: {', '.join(settings.ALLOWED_EXTENSIONS)}"

    # Verificação de Magic Bytes para formatos binários
    if ext in MAGIC_HEADERS:
        is_valid_magic = any(content.startswith(magic) for magic in MAGIC_HEADERS[ext])
        if not is_valid_magic:
            return False, f"Assinatura de arquivo inválida ou corrompida para o formato .{ext}."

    return True, None


# ============================
# Extratores com Mapeamento de Localização
# ============================

def extract_pdf(content: bytes) -> List[ExtractedUnit]:
    units = []
    stream = BytesIO(content)
    reader = PdfReader(stream)
    for index, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            units.append(ExtractedUnit(location=f"Página {index + 1}", text=text.strip()))
    return units


def extract_csv(content: bytes) -> List[ExtractedUnit]:
    units = []
    try:
        decoded = content.decode("utf-8")
    except UnicodeDecodeError:
        decoded = content.decode("latin-1", errors="replace")

    reader = csv.reader(StringIO(decoded))
    headers = []
    for row_idx, row in enumerate(reader, start=1):
        if row_idx == 1 and not headers:
            headers = [h.strip() for h in row]
            units.append(ExtractedUnit(location="Cabeçalho CSV", text=" | ".join(headers)))
            continue
        for col_idx, value in enumerate(row):
            col_name = headers[col_idx] if col_idx < len(headers) and headers[col_idx] else f"Coluna {col_idx + 1}"
            if value.strip():
                units.append(ExtractedUnit(
                    location=f"Coluna {col_name}, Linha {row_idx}",
                    text=value.strip()
                ))
    return units


def extract_xlsx(content: bytes) -> List[ExtractedUnit]:
    units = []
    wb = load_workbook(filename=BytesIO(content), data_only=True, read_only=True)
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row in sheet.iter_rows(values_only=False):
            for cell in row:
                if cell.value is not None:
                    val_str = str(cell.value).strip()
                    if val_str:
                        units.append(ExtractedUnit(
                            location=f"Aba '{sheet_name}', Célula {cell.coordinate}",
                            text=val_str
                        ))
    wb.close()
    return units


def extract_docx(content: bytes) -> List[ExtractedUnit]:
    units = []
    doc = docx.Document(BytesIO(content))
    for p_idx, paragraph in enumerate(doc.paragraphs, start=1):
        if paragraph.text.strip():
            units.append(ExtractedUnit(
                location=f"Parágrafo {p_idx}",
                text=paragraph.text.strip()
            ))

    for t_idx, table in enumerate(doc.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            for c_idx, cell in enumerate(row.cells, start=1):
                if cell.text.strip():
                    units.append(ExtractedUnit(
                        location=f"Tabela {t_idx}, Linha {r_idx}, Coluna {c_idx}",
                        text=cell.text.strip()
                    ))
    return units


def extract_txt_log(content: bytes) -> List[ExtractedUnit]:
    units = []
    try:
        decoded = content.decode("utf-8")
    except UnicodeDecodeError:
        decoded = content.decode("latin-1", errors="replace")

    for line_idx, line in enumerate(decoded.splitlines(), start=1):
        if line.strip():
            units.append(ExtractedUnit(
                location=f"Linha {line_idx}",
                text=line.strip()
            ))
    return units


def extract_json(content: bytes) -> List[ExtractedUnit]:
    units = []
    try:
        data = json.loads(content.decode("utf-8"))
    except UnicodeDecodeError:
        data = json.loads(content.decode("latin-1", errors="replace"))

    def traverse(obj, path=""):
        if isinstance(obj, dict):
            for key, val in obj.items():
                current_path = f"{path}.{key}" if path else key
                traverse(val, current_path)
        elif isinstance(obj, list):
            for idx, item in enumerate(obj):
                current_path = f"{path}[{idx}]"
                traverse(item, current_path)
        else:
            units.append(ExtractedUnit(
                location=f"Campo {path or 'root'}",
                text=str(obj)
            ))

    traverse(data)
    return units


def extract_xml(content: bytes) -> List[ExtractedUnit]:
    units = []
    dom = xml_minidom.parseString(content)
    
    def traverse_node(node, path=""):
        if node.nodeType == node.ELEMENT_NODE:
            current_path = f"{path}/{node.tagName}" if path else node.tagName
            if node.attributes:
                for attr_name, attr_val in node.attributes.items():
                    units.append(ExtractedUnit(
                        location=f"Atributo {current_path}/@{attr_name}",
                        text=attr_val
                    ))
            for child in node.childNodes:
                traverse_node(child, current_path)
        elif node.nodeType == node.TEXT_NODE and node.data.strip():
            units.append(ExtractedUnit(
                location=f"Tag <{path}>",
                text=node.data.strip()
            ))

    traverse_node(dom.documentElement)
    return units


# ============================
# Ponto de Entrada de Leitura
# ============================

def read_file_structural(file_name: str, content: bytes) -> FileExtractionResult:
    """Extrai texto e estruturas retendo informações contextuais de posição."""
    safe_name = sanitize_filename(file_name)
    is_valid, err_msg = validate_file_content(safe_name, content)
    
    if not is_valid:
        return FileExtractionResult(
            success=False,
            file_name=safe_name,
            file_type="unknown",
            file_size=len(content),
            error=err_msg
        )

    ext = safe_name.rsplit(".", 1)[-1].lower()

    try:
        if ext == "pdf":
            units = extract_pdf(content)
        elif ext == "csv":
            units = extract_csv(content)
        elif ext == "xlsx":
            units = extract_xlsx(content)
        elif ext == "docx":
            units = extract_docx(content)
        elif ext in ("txt", "log"):
            units = extract_txt_log(content)
        elif ext == "json":
            units = extract_json(content)
        elif ext == "xml":
            units = extract_xml(content)
        else:
            return FileExtractionResult(
                success=False,
                file_name=safe_name,
                file_type=ext,
                file_size=len(content),
                error="Formato não suportado."
            )

        return FileExtractionResult(
            success=True,
            file_name=safe_name,
            file_type=ext,
            file_size=len(content),
            units=units
        )
    except Exception as e:
        return FileExtractionResult(
            success=False,
            file_name=safe_name,
            file_type=ext,
            file_size=len(content),
            error=f"Falha ao processar conteúdo do arquivo: {str(e)}"
        )
